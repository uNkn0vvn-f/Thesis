"""Stratified samples of fund filings and company-years, with blind coding sheets and keys.

The answer keys carry each item's design weight; cross_model_labels.py reads them to weight
its agreement statistics between the two language models. The coding sheets hide the model
answers so that a human coder could score the same items blind. Strata are mutually exclusive (first matching rule wins) so each item carries a
design weight N_stratum / n_stratum; a value-proportional stratum carries 1 / inclusion
probability. Seeded, so the draws are reproducible.

Outputs: data/interim/validation/ (coding sheets contain filing text and are not part of the repository).
"""
from __future__ import annotations

import gzip
import json
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT))
from src.analysis.robustness.common import (INTERIM, PROCESSED, ROLES,  # noqa: E402
                                                 load_exposure_table, load_holdings)

OUT = INTERIM / 'validation'
OUT.mkdir(parents=True, exist_ok=True)
SEED = 20261011
FUND_N = {'F1 prompt disagreement': 20, 'F2 theme, developer named': 20,
          'F3 theme, enabler named': 20, 'F4 theme, adopter named': 20,
          'F5 theme, no role named': 20, 'F6 process claim only': 20,
          'F7 AI mentioned, no theme or process claim': 20}
FIRM_N = {'C0 ten largest holdings (take-all)': None, 'C1 prompts straddle the role threshold': 25,
          'C2 no filing or patent evidence': 20, 'C3 developer >= 2': 20, 'C4 enabler >= 2': 20,
          'C5 adopter >= 2': 20, 'C6 no role >= 2 (value-proportional)': 25}


def edgar_link(cik, accession):
    if not accession or pd.isna(accession):
        return ''
    return 'https://www.sec.gov/Archives/edgar/data/%s/%s/' % (int(float(cik)), accession.replace('-', ''))


def fund_sample():
    C = pd.read_csv(PROCESSED / 'claim_filings.csv', dtype={'series_id': str})
    ev = {}
    with gzip.open(INTERIM / 'fund_claims' / 'evidence.jsonl.gz', 'rt', encoding='utf-8') as fh:
        for line in fh:
            r = json.loads(line)
            ev[(r['accession'], r['series_id'])] = r
    C['evidence'] = [ev.get((a, s), {}).get('evidence', '') for a, s in zip(C.accession, C.series_id)]
    C['registrant'] = [ev.get((a, s), {}).get('company', '') for a, s in zip(C.accession, C.series_id)]
    C['cik'] = [ev.get((a, s), {}).get('cik', '') for a, s in zip(C.accession, C.series_id)]
    C = C[C.evidence.str.len() > 0].copy()
    E = load_exposure_table('deepseek')  # the sample was drawn on the DeepSeek labels and is fixed
    sample_ids = set(E[E.theme_broad & (E.labelled >= 50)].series_id)
    C['in_principal_sample_fund'] = C.series_id.isin(sample_ids)
    agree = C[['invests_ai_theme_agree', 'uses_ai_agree', 'chain_developer_agree',
               'chain_enabler_agree', 'chain_adopter_agree']].all(axis=1)
    rules = [('F1 prompt disagreement', ~agree),
             ('F2 theme, developer named', C.theme_broad & C.chain_developer),
             ('F3 theme, enabler named', C.theme_broad & C.chain_enabler),
             ('F4 theme, adopter named', C.theme_broad & C.chain_adopter),
             ('F5 theme, no role named', C.theme_broad),
             ('F6 process claim only', C.uses_ai),
             ('F7 AI mentioned, no theme or process claim', pd.Series(True, index=C.index))]
    C['stratum'] = None
    for name, m in rules:
        C.loc[C.stratum.isna() & m, 'stratum'] = name
    rng = np.random.default_rng(SEED)
    picks = []
    for name, n in FUND_N.items():
        pool = C[C.stratum == name]
        take = pool.sample(n=min(n, len(pool)), random_state=int(rng.integers(1 << 31)))
        take = take.assign(N_stratum=len(pool), n_stratum=len(take), weight=len(pool) / len(take))
        picks.append(take)
    S = pd.concat(picks)
    S = S.sample(frac=1, random_state=SEED).reset_index(drop=True)
    S['item_id'] = ['F%03d' % (i + 1) for i in range(len(S))]
    S['edgar_filing_index'] = [edgar_link(c, a) if c else '' for c, a in zip(S.cik, S.accession)]
    return S, C


def firm_sample():
    E = load_exposure_table('deepseek')  # the sample was drawn on the DeepSeek labels and is fixed
    H = load_holdings()
    D = E[E.theme_broad & (E.labelled >= 50)][['series_id', 'q']]
    h = H.merge(D, on=['series_id', 'q'])
    val = h.groupby(['gvkey', 'yr']).value.sum().rename('value_held').reset_index()
    rows = []
    for y in sorted(val.yr.unique()):
        v = pd.read_parquet(INTERIM / 'firm_labels' / 'vintages' / '02_label' / str(y) / 'firm_ai_labels.parquet')
        v['gvkey'] = v.gvkey.astype(str)
        v['yr'] = y
        rows.append(v)
    L = pd.concat(rows, ignore_index=True)
    F = val.merge(L, on=['gvkey', 'yr'], how='inner')
    straddle = pd.Series(False, index=F.index)
    for r in ROLES:
        straddle |= (F[r + '_1'] >= 2) != (F[r + '_2'] >= 2)
    top = F.sort_values('value_held', ascending=False).drop_duplicates('gvkey').head(10).index
    rules = [('C0 ten largest holdings (take-all)', F.index.isin(top)),
             ('C1 prompts straddle the role threshold', straddle),
             ('C2 no filing or patent evidence', ~F.has_evidence.astype(bool)),
             ('C3 developer >= 2', F.developer >= 2), ('C4 enabler >= 2', F.enabler >= 2),
             ('C5 adopter >= 2', F.adopter >= 2),
             ('C6 no role >= 2 (value-proportional)', pd.Series(True, index=F.index))]
    F['stratum'] = None
    for name, m in rules:
        F.loc[F.stratum.isna() & np.asarray(m), 'stratum'] = name
    rng = np.random.default_rng(SEED + 1)
    picks = []
    for name, n in FIRM_N.items():
        pool = F[F.stratum == name]
        if n is None:
            take = pool.assign(N_stratum=len(pool), n_stratum=len(pool), weight=1.0)
        elif name.startswith('C6'):
            p = pool.value_held / pool.value_held.sum()
            idx = rng.choice(pool.index, size=min(n, len(pool)), replace=False, p=p.values)
            # approximate inclusion probability for a without-replacement PPS draw
            pi = np.minimum(1.0, n * p.loc[idx])
            take = pool.loc[idx].assign(N_stratum=len(pool), n_stratum=len(idx), weight=1 / pi.values)
        else:
            take = pool.sample(n=min(n, len(pool)), random_state=int(rng.integers(1 << 31)))
            take = take.assign(N_stratum=len(pool), n_stratum=len(take), weight=len(pool) / len(take))
        picks.append(take)
    S = pd.concat(picks)
    # evidence text exactly as available for that vintage
    texts = {}
    for y in sorted(S.yr.unique()):
        p = INTERIM / 'firm_labels' / 'vintages' / '00_evidence' / str(y) / 'business_text.jsonl'
        if p.exists():
            with open(p, encoding='utf-8', errors='replace') as fh:
                for line in fh:
                    try:
                        x = json.loads(line)
                    except json.JSONDecodeError:
                        continue
                    texts[(str(x['gvkey']), y)] = (x.get('business_text') or '', x.get('acc'), x.get('cik'))
        pp = INTERIM / 'firm_labels' / 'vintages' / '00_evidence' / str(y) / 'patent_profile.parquet'
        if pp.exists():
            P = pd.read_parquet(pp)
            P['gvkey'] = P.gvkey.astype(str)
            for g, prof in zip(P.gvkey, P.cpc_profile):
                texts.setdefault(('pat', g, y), prof)
    S['business_text_excerpt'] = [texts.get((g, y), ('',))[0][:3500] or 'NO ANNUAL-REPORT TEXT SUPPLIED TO THE MODEL'
                                  for g, y in zip(S.gvkey, S.yr)]
    S['patent_profile'] = [texts.get(('pat', g, y), '') or '' for g, y in zip(S.gvkey, S.yr)]
    S['edgar_filing_index'] = [edgar_link(texts.get((g, y), ('', None, None))[2], texts.get((g, y), ('', None, None))[1])
                               if texts.get((g, y), ('', None, None))[1] else '' for g, y in zip(S.gvkey, S.yr)]
    S = S.sample(frac=1, random_state=SEED).reset_index(drop=True)
    S['item_id'] = ['C%03d' % (i + 1) for i in range(len(S))]
    return S, F


def write_sheet(path, df, show, code_cols, validations, widths):
    wb = Workbook()
    ws = wb.active
    ws.title = 'coding'
    cols = show + code_cols
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill('solid', fgColor='DDDDDD')
    for _, r in df.iterrows():
        vals = []
        for c in cols:
            v = r.get(c, '') if c in show else ''
            vals.append('' if v is None or (not isinstance(v, str) and pd.isna(v)) else
                        (v.item() if hasattr(v, 'item') else v))
        ws.append(vals)
    for i, c in enumerate(cols, start=1):
        letter = ws.cell(row=1, column=i).column_letter
        ws.column_dimensions[letter].width = widths.get(c, 14)
        for cell in ws[letter][1:]:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            if c in code_cols:
                cell.fill = PatternFill('solid', fgColor='FFF7D6')
    for col, options in validations.items():
        letter = ws.cell(row=1, column=cols.index(col) + 1).column_letter
        dv = DataValidation(type='list', formula1='"%s"' % ','.join(options), allow_blank=True)
        ws.add_data_validation(dv)
        dv.add('%s2:%s%d' % (letter, letter, len(df) + 1))
    ws.freeze_panes = 'C2'
    wb.save(path)


def main():
    FS, FC = fund_sample()
    show = ['item_id', 'series_name', 'registrant', 'date', 'accession', 'edgar_filing_index', 'evidence']
    yn = ['Y', 'N']
    code = ['Q1_ai_theme', 'Q2_ai_selection_basis', 'Q3_developer', 'Q4_enabler', 'Q5_adopter',
            'Q6_uses_ai_in_decisions', 'Q7_risk_only', 'Q8_text_belongs_to_this_fund',
            'confidence_1to3', 'supporting_quote', 'notes', 'coder_initials']
    write_sheet(OUT / 'fund_claims_coding_sheet.xlsx', FS, show, code,
                {**{c: yn for c in code[:7]}, 'Q8_text_belongs_to_this_fund': ['Y', 'N', 'unsure'],
                 'confidence_1to3': ['1', '2', '3']},
                {'evidence': 90, 'series_name': 28, 'registrant': 22, 'edgar_filing_index': 30,
                 'supporting_quote': 40, 'notes': 30})
    keyf = ['item_id', 'stratum', 'N_stratum', 'n_stratum', 'weight', 'series_id', 'accession',
            'in_principal_sample_fund', 'theme_broad', 'theme_strict', 'uses_ai', 'chain_developer',
            'chain_enabler', 'chain_adopter', 'strategy_text_risk_only', 'invests_ai_theme_agree',
            'uses_ai_agree', 'chain_developer_agree', 'chain_enabler_agree', 'chain_adopter_agree']
    FS[keyf].to_csv(OUT / 'KEY_fund_claims_model_labels_DO_NOT_OPEN_BEFORE_CODING.csv', index=False)
    FC.groupby('stratum').size().rename('N').to_csv(OUT / 'frame_fund_claims_strata.csv')

    CS, CF = firm_sample()
    show = ['item_id', 'name', 'ctry', 'sic4', 'yr', 'edgar_filing_index', 'business_text_excerpt', 'patent_profile']
    code = ['developer_0to3', 'enabler_0to3', 'adopter_0to3', 'ai_revenue_share_band',
            'evidence_sufficient', 'would_change_with_outside_knowledge', 'confidence_1to3', 'notes',
            'coder_initials']
    sc = ['0', '1', '2', '3']
    write_sheet(OUT / 'company_roles_coding_sheet.xlsx', CS, show, code,
                {'developer_0to3': sc, 'enabler_0to3': sc, 'adopter_0to3': sc,
                 'ai_revenue_share_band': ['0', '1-10', '10-25', '25-50', '50+', 'cannot tell'],
                 'evidence_sufficient': yn, 'would_change_with_outside_knowledge': yn,
                 'confidence_1to3': ['1', '2', '3']},
                {'business_text_excerpt': 90, 'patent_profile': 40, 'name': 26, 'edgar_filing_index': 30,
                 'notes': 30})
    keyc = ['item_id', 'stratum', 'N_stratum', 'n_stratum', 'weight', 'gvkey', 'yr', 'value_held',
            'has_text', 'has_pat', 'has_evidence', 'developer', 'enabler', 'adopter', 'developer_1',
            'developer_2', 'enabler_1', 'enabler_2', 'adopter_1', 'adopter_2', 'ai_rev', 'business_1',
            'sales_basis_1']
    CS[keyc].to_csv(OUT / 'KEY_company_roles_model_labels_DO_NOT_OPEN_BEFORE_CODING.csv', index=False)
    CF.groupby('stratum').agg(N=('gvkey', 'size'), value_held=('value_held', 'sum')).to_csv(
        OUT / 'frame_company_roles_strata.csv')
    summary = {'fund_items': len(FS), 'company_items': len(CS),
               'fund_strata': FS.groupby('stratum').size().to_dict(),
               'company_strata': CS.groupby('stratum').size().to_dict(),
               'fund_items_from_principal_sample_funds': int(FS.in_principal_sample_fund.sum()), 'seed': SEED}
    (OUT / 'sampling_summary.json').write_text(json.dumps(summary, indent=2), encoding='utf-8')
    print(json.dumps(summary, indent=2))


if __name__ == '__main__':
    main()
